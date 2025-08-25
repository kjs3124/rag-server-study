from typing import List, Any, Dict, Optional
import logging
from docx import Document

from .base import BaseDocumentParser, ParsedDocument, DocumentChunk

logger = logging.getLogger(__name__)

class DOCXParser(BaseDocumentParser):
    """DOCX 문서 파서 - python-docx + LangChain 청킹"""
    
    def parse(self, file_path: str, chunk_size: int = 1000, **kwargs) -> ParsedDocument:
        """DOCX 파일을 파싱하여 청크로 분할"""
        
        logger.info(f"Using python-docx + LangChain for DOCX parsing: {file_path}")
        return self._parse_with_python_docx(file_path, chunk_size, **kwargs)
    
    
    def _parse_with_python_docx(self, file_path: str, chunk_size: int, **kwargs) -> ParsedDocument:
        """python-docx + LangChain을 사용한 최적화된 DOCX 파싱"""
        
        doc = Document(file_path)
        
        # 전체 텍스트 추출
        full_text_parts: List[str] = []
        current_section = None
        
        # 문서 메타데이터 추출
        doc_properties = {
            'title': doc.core_properties.title or '',
            'author': doc.core_properties.author or '',
            'subject': doc.core_properties.subject or '',
            'created': str(doc.core_properties.created) if doc.core_properties.created else '',
            'modified': str(doc.core_properties.modified) if doc.core_properties.modified else ''
        }
        
        # 문단별로 텍스트 추출
        for para in doc.paragraphs:
            if not para.text.strip():
                continue
                
            # 스타일 정보 추출
            style_name = para.style.name if para.style else 'Normal'
            is_heading = (style_name and 
                        (style_name.startswith('Heading') or style_name.startswith('Title')))
            
            # 헤더인 경우 섹션 구분 추가
            if is_heading:
                current_section = para.text.strip()
                full_text_parts.append(f"\n\n# {para.text.strip()}\n")
            else:
                full_text_parts.append(para.text)
        
        # 테이블 텍스트 추출 및 추가
        for table_idx, table in enumerate(doc.tables):
            table_text = [f"\n\nTable {table_idx + 1}:"]
            
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    row_text.append(cell_text)
                if row_text:
                    table_text.append(' | '.join(row_text))
            
            if len(table_text) > 1:
                full_text_parts.append('\n'.join(table_text))
        
        # 전체 텍스트 결합
        full_text = '\n'.join(full_text_parts)
        
        # LangChain 청킹 (base 클래스 메서드 사용)
        chunks = self._create_langchain_chunks(
            full_text,
            chunk_size,
            file_path,
            separators=[
                "\n\n# ",    # 헤딩 구분
                "\n\n",      # 문단 구분
                "\nTable",   # 테이블 구분
                "\n",        # 줄 구분
                ". ",        # 문장 구분
                " ",         # 공백 구분
                ""           # 문자 구분
            ],
            **doc_properties,
            parser="python_docx_langchain",
            has_tables=len(doc.tables) > 0,
            table_count=len(doc.tables)
        )
        
        return ParsedDocument(
            chunks=chunks,
            metadata={
                **self._extract_metadata(file_path),
                **doc_properties,
                "parser": "python_docx_langchain",
                "total_chunks": len(chunks),
                "has_tables": len(doc.tables) > 0,
                "table_count": len(doc.tables)
            },
            file_type="docx"
        )
    
    
    def get_supported_extensions(self) -> List[str]:
        return ['.docx', '.doc']