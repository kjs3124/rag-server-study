from typing import List
from pptx import Presentation

from .base import BaseDocumentParser, ParsedDocument, DocumentChunk

class PPTXParser(BaseDocumentParser):
    """PPTX 파서 - python-pptx + LangChain 청킹"""
    
    def parse(self, file_path: str, chunk_size: int = 1000, **kwargs) -> ParsedDocument:
        """PPTX 파일을 파싱하여 청크로 분할"""
        
        prs = Presentation(file_path)
        
        # 전체 텍스트 추출
        all_text_parts: List[str] = []
        slide_info_list: List[dict] = []
        
        for slide_idx, slide in enumerate(prs.slides):
            slide_text_parts: List[str] = []
            
            # 슬라이드 번호 헤더 추가
            slide_text_parts.append(f"\n\nSlide {slide_idx + 1}:")
            
            for shape in slide.shapes:
                if hasattr(shape, "text_frame") and shape.text_frame:
                    text = shape.text_frame.text.strip()
                    if text:
                        slide_text_parts.append(text)
                elif hasattr(shape, "text"):
                    text = shape.text.strip()
                    if text:
                        slide_text_parts.append(text)
            
            if len(slide_text_parts) > 1:  # 헤더 외에 텍스트가 있으면
                slide_content = '\n'.join(slide_text_parts)
                all_text_parts.append(slide_content)
                
                slide_info_list.append({
                    'slide_number': slide_idx + 1,
                    'layout_name': slide.slide_layout.name if hasattr(slide.slide_layout, 'name') else 'Unknown',
                    'text_count': len(slide_text_parts) - 1  # 헤더 제외
                })
        
        # 전체 텍스트 결합
        full_text = '\n'.join(all_text_parts)
        
        # LangChain 청킹 (base 클래스 메서드 사용)
        chunks = self._create_langchain_chunks(
            full_text,
            chunk_size,
            file_path,
            separators=[
                "\n\nSlide ",   # 슬라이드 구분
                "\n\n",        # 문단 구분
                "\n",          # 줄 구분
                ". ",          # 문장 구분
                " ",           # 공백 구분
                ""             # 문자 구분
            ],
            parser="pptx_langchain",
            slides_with_text=len(slide_info_list)
        )
        
        # 슬라이드 번호 추가 (후처리)
        for chunk in chunks:
            slide_number = self._extract_slide_number(chunk.content)
            chunk.metadata["slide_number"] = slide_number
            chunk.metadata["is_slide_content"] = "Slide " in chunk.content
            chunk.page_number = slide_number
        
        return ParsedDocument(
            chunks=chunks,
            metadata={
                **self._extract_metadata(file_path), 
                "parser": "pptx_langchain", 
                "total_slides": len(prs.slides),
                "slides_with_text": len(slide_info_list),
                "total_chunks": len(chunks)
            },
            file_type="pptx"
        )
    
    
    def _extract_slide_number(self, text: str) -> int:
        """텍스트에서 슬라이드 번호 추출"""
        import re
        match = re.search(r'Slide (\d+):', text)
        if match:
            return int(match.group(1))
        return 1
    
    def get_supported_extensions(self) -> List[str]:
        return ['.pptx', '.ppt']