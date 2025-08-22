from typing import List, Any, Dict
import zipfile
import xml.etree.ElementTree as ET

try:
    from docx import Document
    PYTHON_DOCX_AVAILABLE = True
except ImportError:
    PYTHON_DOCX_AVAILABLE = False

from .base import BaseDocumentParser, ParsedDocument, DocumentChunk

class DOCXParser(BaseDocumentParser):
    """DOCX 문서 파서"""
    
    def parse(self, file_path: str, chunk_size: int = 1000, **kwargs) -> ParsedDocument:
        """DOCX 파일을 파싱하여 청크로 분할"""
        
        if PYTHON_DOCX_AVAILABLE:
            return self._parse_with_python_docx(file_path, chunk_size, **kwargs)
        else:
            return self._parse_with_xml(file_path, chunk_size, **kwargs)
    
    def _parse_with_python_docx(self, file_path: str, chunk_size: int, **kwargs) -> ParsedDocument:
        """python-docx를 사용한 고급 DOCX 파싱"""
        try:
            doc = Document(file_path)
            
            chunks: List[DocumentChunk] = []
            current_chunk: List[Any] = []
            current_size = 0
            
            # 문단별로 처리
            for para in doc.paragraphs:
                if not para.text.strip():
                    continue
                
                # 스타일 정보 추출
                style_name = para.style.name if para.style else 'Normal'
                
                para_text = para.text
                para_size = len(para_text)
                
                # 청크 크기 초과시 분할
                if current_size + para_size > chunk_size and current_chunk:
                    chunk_content = '\n'.join([p['text'] for p in current_chunk])
                    chunks.append(self._create_chunk_with_metadata(
                        chunk_content,
                        file_path,
                        len(chunks),
                        current_chunk
                    ))
                    current_chunk = []
                    current_size = 0
                
                current_chunk.append({
                    'text': para_text,
                    'style': style_name,
                    'is_heading': style_name.startswith('Heading')
                })
                current_size += para_size
            
            # 남은 청크 처리
            if current_chunk:
                chunk_content = '\n'.join([p['text'] for p in current_chunk])
                chunks.append(self._create_chunk_with_metadata(
                    chunk_content,
                    file_path,
                    len(chunks),
                    current_chunk
                ))
            
            # 테이블 처리
            table_chunks = self._extract_tables(doc, file_path, len(chunks))
            chunks.extend(table_chunks)
            
            return ParsedDocument(
                chunks=chunks,
                metadata={
                    **self._extract_metadata(file_path),
                    "parser": "python-docx",
                    "total_chunks": len(chunks),
                    "has_tables": len(table_chunks) > 0
                },
                file_type="docx"
            )
            
        except Exception as e:
            raise Exception(f"DOCX 파싱 중 오류 발생: {str(e)}")
    
    def _parse_with_xml(self, file_path: str, chunk_size: int, **kwargs) -> ParsedDocument:
        """XML을 직접 파싱하는 방법 (fallback)"""
        try:
            chunks: List[DocumentChunk] = []
            
            with zipfile.ZipFile(file_path, 'r') as docx:
                # document.xml 파일 읽기
                xml_content = docx.read('word/document.xml')
                root = ET.fromstring(xml_content)
                
                # 네임스페이스 정의
                namespaces = {
                    'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
                }
                
                # 문단 텍스트 추출
                paragraphs = root.findall('.//w:p', namespaces)
                
                all_text = []
                for para in paragraphs:
                    para_text = ''
                    for text_elem in para.findall('.//w:t', namespaces):
                        if text_elem.text:
                            para_text += text_elem.text
                    
                    if para_text.strip():
                        all_text.append(para_text)
                
                # 텍스트를 청크로 분할
                full_text = '\n'.join(all_text)
                chunks = self._split_text_into_chunks(full_text, file_path, chunk_size)
                
                return ParsedDocument(
                    chunks=chunks,
                    metadata={
                        **self._extract_metadata(file_path),
                        "parser": "xml_parser",
                        "total_chunks": len(chunks)
                    },
                    file_type="docx"
                )
                
        except Exception as e:
            raise Exception(f"DOCX XML 파싱 중 오류 발생: {str(e)}")
    
    def _extract_tables(self, doc, file_path: str, start_index: int) -> List[DocumentChunk]:
        """문서에서 테이블 추출"""
        table_chunks: List[DocumentChunk] = []
        
        for table_idx, table in enumerate(doc.tables):
            table_text = []
            
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    row_text.append(cell_text)
                table_text.append(' | '.join(row_text))
            
            if table_text:
                table_content = '\n'.join(table_text)
                chunk = DocumentChunk(
                    content=table_content,
                    metadata={
                        **self._extract_metadata(file_path),
                        "content_type": "table",
                        "table_index": table_idx,
                        "row_count": len(table.rows),
                        "column_count": len(table.columns)
                    },
                    chunk_id=self._create_chunk_id(file_path, start_index + table_idx)
                )
                table_chunks.append(chunk)
        
        return table_chunks
    
    def _create_chunk_with_metadata(self, content: str, file_path: str, 
                                   chunk_index: int, para_info: List[dict]) -> DocumentChunk:
        """메타데이터가 포함된 청크 생성"""
        
        # 헤딩 정보 추출
        headings = [p['text'] for p in para_info if p.get('is_heading', False)]
        styles = list(set(p['style'] for p in para_info))
        
        return DocumentChunk(
            content=content,
            metadata={
                **self._extract_metadata(file_path),
                "chunk_index": chunk_index,
                "character_count": len(content),
                "paragraph_count": len(para_info),
                "headings": headings,
                "styles_used": styles,
                "has_headings": len(headings) > 0
            },
            chunk_id=self._create_chunk_id(file_path, chunk_index),
            section_title=headings[0] if headings else None
        )
    
    def _split_text_into_chunks(self, text: str, file_path: str, chunk_size: int) -> List[DocumentChunk]:
        """텍스트를 청크로 분할 (fallback)"""
        chunks: List[DocumentChunk] = []
        words = text.split()
        
        current_chunk: List[Any] = []
        current_size = 0
        
        for word in words:
            word_size = len(word) + 1
            if current_size + word_size > chunk_size and current_chunk:
                chunk_text = ' '.join(current_chunk)
                chunk = DocumentChunk(
                    content=chunk_text,
                    metadata={
                        **self._extract_metadata(file_path),
                        "chunk_index": len(chunks),
                        "character_count": len(chunk_text)
                    },
                    chunk_id=self._create_chunk_id(file_path, len(chunks))
                )
                chunks.append(chunk)
                current_chunk = [word]
                current_size = word_size
            else:
                current_chunk.append(word)
                current_size += word_size
        
        if current_chunk:
            chunk_text = ' '.join(current_chunk)
            chunk = DocumentChunk(
                content=chunk_text,
                metadata={
                    **self._extract_metadata(file_path),
                    "chunk_index": len(chunks),
                    "character_count": len(chunk_text)
                },
                chunk_id=self._create_chunk_id(file_path, len(chunks))
            )
            chunks.append(chunk)
        
        return chunks
    
    def get_supported_extensions(self) -> List[str]:
        return ['.docx', '.doc']